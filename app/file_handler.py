import pandas as pd
import pdfplumber
from docx import Document
from pptx import Presentation
from bs4 import BeautifulSoup
import json 
import xml.etree.ElementTree as ET
from io import BytesIO

def extract_file_data(filename, file_bytes):
    ext = filename.split('.')[-1].lower()
    buffer = BytesIO(file_bytes)

    if ext == "csv":
        return pd.read_csv(buffer)
    elif ext in ["xls", "xlsx"]:
        return pd.read_excel(buffer)
    elif ext == "pdf":
        with pdfplumber.open(buffer) as pdf:
            text = "\n".join(page.extract_text() for page in pdf.pages if page.extract_text())
        return pd.DataFrame({"text": text.split('\n')})
    elif ext == "docx":
        doc = Document(buffer)
        lines = [para.text for para in doc.paragraphs if para.text.strip()]
        return pd.DataFrame({"text": lines})
    elif ext == "txt":
        text = buffer.read().decode("utf-8").splitlines()
        return pd.DataFrame({"text": text})
    elif ext == "pptx":
        prs = Presentation(buffer)
        lines = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    lines.append(shape.text)
        return pd.DataFrame({"text": lines})
    elif ext == "json":
        content = json.load(buffer)
        return pd.json_normalize(content)
    elif ext == "xml":
        tree = ET.parse(buffer)
        root = tree.getroot()
        records = [{elem.tag: elem.text for elem in child} for child in root]
        return pd.DataFrame(records)
    elif ext == "html":
        soup = BeautifulSoup(buffer.read(), "html.parser")
        text = [t.strip() for t in  soup.stripped_strings]
        return pd.DataFrame({"text": text})
    else:
         raise ValueError(f"Unsupported File Format: {ext}")